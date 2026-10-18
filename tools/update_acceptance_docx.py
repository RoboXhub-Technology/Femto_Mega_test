from __future__ import annotations

import hashlib
import os
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


DOCX_PATH = Path(r"D:\Master\Orbbec_Femto_Mega_3D深度相机验收记录表.docx")


def media_hashes(path: Path) -> dict[str, str]:
    with zipfile.ZipFile(path) as archive:
        return {
            name: hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist()
            if name.startswith("word/media/") and not name.endswith("/")
        }


def add_result(cell, lines: list[str]) -> None:
    marker = "【本机实测结果（2026-10-18）】"
    for index, existing in enumerate(cell.paragraphs):
        if marker in existing.text:
            if len(lines) == 1:
                following = cell.paragraphs[index + 1] if index + 1 < len(cell.paragraphs) else None
                existing.clear()
                marker_run = existing.add_run(marker + " ")
                marker_run.bold = True
                marker_run.font.size = Pt(8.5)
                marker_run.font.color.rgb = RGBColor(31, 78, 121)
                marker_run.font.name = "Microsoft YaHei"
                marker_run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
                body_run = existing.add_run(lines[0])
                body_run.font.size = Pt(8)
                body_run.font.name = "Microsoft YaHei"
                body_run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
                if following is not None:
                    following._element.getparent().remove(following._element)
                existing.paragraph_format.keep_together = True
            existing.paragraph_format.keep_with_next = True
            return

    heading = cell.add_paragraph()
    heading.paragraph_format.space_before = Pt(2)
    heading.paragraph_format.space_after = Pt(0)
    heading.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    heading.paragraph_format.keep_with_next = True
    run = heading.add_run(marker)
    run.bold = True
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(31, 78, 121)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    if len(lines) == 1:
        heading.paragraph_format.keep_together = True
        body_run = heading.add_run(" " + lines[0])
        body_run.font.size = Pt(8)
        body_run.font.color.rgb = RGBColor(0, 0, 0)
        body_run.font.name = "Microsoft YaHei"
        body_run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        return

    for text in lines:
        paragraph = cell.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        run = paragraph.add_run(text)
        run.font.size = Pt(8)
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def main() -> int:
    if not DOCX_PATH.exists():
        raise FileNotFoundError(DOCX_PATH)

    before_media = media_hashes(DOCX_PATH)
    document = Document(DOCX_PATH)
    if len(document.tables) != 1 or len(document.tables[0].rows) < 13:
        raise RuntimeError("Unexpected acceptance-form structure; refusing to edit.")

    table = document.tables[0]
    additions: dict[int, list[str]] = {
        2: [
            "设备：Orbbec Femto Mega；序列号 CL3N241005X；固件 1.3.0；设备枚举/打开正常。",
            "pyorbbecsdk 运行时报告 SDK 2.9.3；K4A Wrapper v2.0.12，使用其配套 OrbbecSDK.dll 2.5.5。",
        ],
        3: [
            "由内参按针孔模型近似：NFOV H64.77°、V59.28°；WFOV H90.84°、V90.84°。",
            "该计算未展开广角畸变后的边缘视场，不能据此对 H75°/V65°、H120°/V120°作最终否决；需用完整畸变模型或物理角度标定复核。当前：待复核。",
        ],
        4: [
            "pyorbbecsdk 列表中要求的 14 个深度组合全部存在，并均已完成约 60 s 取流；K4A Wrapper full 测试中 14 项均取得 5/5 有效帧且时间戳单调。",
            "结论：所列深度分辨率@帧率组合通过。证据：evidence-python/01～14；k4a_wrapper_compat_test/results/reference_CL3N241005X_2026-10-18/k4a_wrapper_full.json。",
        ],
        5: [
            "未完成 0.25～5.46 m 标定距离实测；未判定。",
        ],
        6: [
            "当前证据未完成 1 m 标准板测距与 0.15% 相对误差计算；本项未判定。",
        ],
        7: [
            "3840×2160@5/15/25 fps 可列出并取流；K4A 三项均为 5/5 PASS。3840×2160@30 fps 未列出，pyorbbecsdk 30 fps 启动失败（退出码 2）。",
            "结论：不满足“3840×2160@≥30 fps”。",
        ],
        8: [
            "3840×2160 profile 内参针孔近似 FOV：H81.17°、V51.45°。结论：满足 H80°、V51°。",
        ],
        9: [
            "OB_PROP_COLOR_HDR_BOOL 未出现在支持属性列表，read_supported=false，write_supported=false；K4A Wrapper 对 Femto Mega 跳过 HDR。",
            "结论：当前设备/固件/SDK 组合未证明支持 RGB HDR，本项未通过。",
        ],
        10: [
            "当前未接入外部校准功率计，未取得最大功耗数据；本项未判定。",
        ],
        11: [
            "Orbbec SDK/pyorbbecsdk：设备、流配置和 IMU 可正常访问。K4A Wrapper v2.0.12：31/31 个 required 状态与基准一致，overall_pass=true。",
            "已覆盖枚举、打开、标定、transformation、RGB 控制项、同步流、14 个深度组合、4K 5/15/25 和 IMU；不含人体骨架识别、HDR、4K30。",
        ],
        12: [
            "pyorbbecsdk 60 s：加速度计 2984 帧、陀螺仪 2984 帧，实测均约 49.777 Hz，帧序号缺口 0、超时 0；K4A IMU 5/5 PASS。",
            "结论：支持 IMU。原有图片保留在本单元格上方。",
        ],
    }

    for row_index, lines in additions.items():
        add_result(table.rows[row_index].cells[1], lines)

    temp_path = DOCX_PATH.with_name(DOCX_PATH.stem + ".updating.docx")
    document.save(temp_path)
    after_media = media_hashes(temp_path)
    if before_media != after_media:
        temp_path.unlink(missing_ok=True)
        raise RuntimeError("Embedded media changed; refusing to replace the original document.")

    os.replace(temp_path, DOCX_PATH)
    print(f"Updated: {DOCX_PATH}")
    print(f"Preserved embedded media files: {len(after_media)}")
    print(f"Tables: {len(document.tables)}; rows in acceptance table: {len(table.rows)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
